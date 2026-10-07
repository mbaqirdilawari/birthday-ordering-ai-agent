"""Excel version of the business case, with live formulas (for colleagues who do not code).

Terminal:
    python -m business_case.excel_model
Creates excel/Birthday_Trike_Business_Case.xlsx
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName

from .model import ASSUMPTIONS, sensitivity

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "excel" / "Birthday_Trike_Business_Case.xlsx"
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT)
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
NOTE = Font(name=FONT, italic=True, color="595959", size=9)
WHITE = Font(name=FONT, bold=True, color="FFFFFF")
HEAD = PatternFill("solid", fgColor="1F6F5C")
INPUT = PatternFill("solid", fgColor="FFF2CC")
DISCLAIMER = "All numbers are SIMULATED assumptions for illustration. They are not real company figures."

LABELS = {
    "cost_social_media_and_influencers": ("Launch cost: social media and influencer posts (PKR)", "#,##0"),
    "cost_posters_and_flyers": ("Launch cost: posters and flyers (PKR)", "#,##0"),
    "cost_trial_events": ("Launch cost: trial events (PKR)", "#,##0"),
    "monthly_event_supervision": ("Running cost: event supervision (PKR per month)", "#,##0"),
    "avg_order_value": ("Average order value (PKR)", "#,##0"),
    "gross_margin": ("Gross margin", "0%"),
    "orders_month_1": ("Orders in month 1", "0"),
    "orders_month_2": ("Orders in month 2", "0"),
    "orders_per_month_after": ("Orders per month from month 3", "0"),
    "seller_earnings_normal_day": ("Trike seller earnings, normal day (PKR)", "#,##0"),
    "seller_earnings_event_day": ("Trike seller earnings, day with an event (PKR)", "#,##0"),
}
NAMES = {
    "cost_social_media_and_influencers": "CostSocial", "cost_posters_and_flyers": "CostPrint",
    "cost_trial_events": "CostTrials", "monthly_event_supervision": "MonthlyCost", "avg_order_value": "AOV",
    "gross_margin": "Margin", "orders_month_1": "Orders1", "orders_month_2": "Orders2",
    "orders_per_month_after": "OrdersSteady", "seller_earnings_normal_day": "EarnNormal",
    "seller_earnings_event_day": "EarnEvent",
}

# Payback month by formula (same logic as model.payback_month_formula), for given steady orders and order value
PAYBACK = (
    'IF(Orders1*{v}*Margin>=LaunchCost+MonthlyCost,1,'
    'IF((Orders1+Orders2)*{v}*Margin>=LaunchCost+2*MonthlyCost,2,'
    'IF({o}*{v}*Margin<=MonthlyCost,"Never",'
    'MAX(3,ROUNDUP((LaunchCost-(Orders1+Orders2)*{v}*Margin+2*{o}*{v}*Margin)/({o}*{v}*Margin-MonthlyCost)-0.000000001,0)))))'
)


def _name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def build(path: Path = OUT) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 16
    ws["A1"], ws["A1"].font = "Birthday Trike Service: Business Case", TITLE
    ws["A2"], ws["A2"].font = DISCLAIMER, NOTE
    ws["A3"], ws["A3"].font = "Edit the blue cells. Everything else recalculates.", NOTE
    r = 5
    for key, (label, fmt) in LABELS.items():
        ws.cell(row=r, column=1, value=label).font = BLACK
        c = ws.cell(row=r, column=2, value=ASSUMPTIONS[key])
        c.font, c.fill, c.number_format = BLUE, INPUT, fmt
        _name(wb, NAMES[key], f"Inputs!$B${r}")
        r += 1
    ws.cell(row=r + 1, column=1, value="Total launch cost (PKR)").font = BOLD
    c = ws.cell(row=r + 1, column=2, value="=CostSocial+CostPrint+CostTrials")
    c.number_format, c.font = "#,##0", BOLD
    _name(wb, "LaunchCost", f"Inputs!$B${r + 1}")

    # Monthly model
    m = wb.create_sheet("Monthly")
    heads = ["Month", "Orders", "Revenue (PKR)", "Gross margin (PKR)", "Cumulative revenue", "Cumulative margin",
             "Cumulative cost (launch + running)", "Net position"]
    for i, h in enumerate(heads, start=1):
        c = m.cell(row=1, column=i, value=h)
        c.font, c.fill, c.alignment = WHITE, HEAD, Alignment(wrap_text=True, horizontal="center")
        m.column_dimensions[chr(64 + i)].width = 18
    months = ASSUMPTIONS["months"]
    for k in range(1, months + 1):
        row = k + 1
        m.cell(row=row, column=1, value=k)
        m.cell(row=row, column=2, value=f"=IF(A{row}=1,Orders1,IF(A{row}=2,Orders2,OrdersSteady))")
        m.cell(row=row, column=3, value=f"=B{row}*AOV")
        m.cell(row=row, column=4, value=f"=C{row}*Margin")
        m.cell(row=row, column=5, value=f"=SUM($C$2:C{row})")
        m.cell(row=row, column=6, value=f"=SUM($D$2:D{row})")
        m.cell(row=row, column=7, value=f"=LaunchCost+MonthlyCost*A{row}")
        m.cell(row=row, column=8, value=f"=F{row}-G{row}")
        for col in range(1, 9):
            cell = m.cell(row=row, column=col)
            cell.font = BLACK
            if col >= 3:
                cell.number_format = "#,##0;(#,##0)"
    m.freeze_panes = "A2"
    last = months + 1

    # Summary
    s = wb.create_sheet("Summary", 0)
    s.column_dimensions["A"].width = 60
    s.column_dimensions["B"].width = 18
    s["A1"], s["A1"].font = "Summary", TITLE
    s["A2"], s["A2"].font = DISCLAIMER, NOTE
    rows = [
        ("Total launch cost (PKR)", "=LaunchCost", "#,##0"),
        ("Orders to recover launch cost (revenue basis, quick pitch method)", "=ROUNDUP(LaunchCost/AOV,0)", "#,##0"),
        ("Month revenue first covers launch cost", f'=IFERROR(INDEX(Monthly!A2:A{last},MATCH(TRUE,INDEX(Monthly!E2:E{last}>=LaunchCost,0),0)),"After month {months}")', "0"),
        ("PAYBACK MONTH (margin basis, including running costs)", "=" + PAYBACK.format(o="OrdersSteady", v="AOV"), "0"),
        ("Trike seller earnings increase on an event day (PKR)", "=EarnEvent-EarnNormal", "#,##0"),
        ("Trike seller earnings increase on an event day (percent)", "=EarnEvent/EarnNormal-1", "0%"),
    ]
    for i, (label, f, fmt) in enumerate(rows, start=4):
        s.cell(row=i, column=1, value=label).font = BOLD if label.isupper() or label.startswith("PAYBACK") else BLACK
        c = s.cell(row=i, column=2, value=f)
        c.number_format, c.font = fmt, BOLD if label.startswith("PAYBACK") else BLACK

    # Sensitivity grid (every cell is a formula)
    g = wb.create_sheet("Sensitivity")
    g["A1"], g["A1"].font = "Payback month (margin basis) by orders per month and average order value", BOLD
    g["A2"], g["A2"].font = DISCLAIMER, NOTE
    grid = sensitivity()
    g.cell(row=4, column=1, value="Orders per month \\ Order value (PKR)").font = WHITE
    g.cell(row=4, column=1).fill = HEAD
    g.column_dimensions["A"].width = 34
    for j, v in enumerate(grid.columns, start=2):
        c = g.cell(row=4, column=j, value=int(v))
        c.font, c.fill, c.number_format = WHITE, HEAD, "#,##0"
        g.column_dimensions[chr(64 + j)].width = 12
    for i, o in enumerate(grid.index, start=5):
        g.cell(row=i, column=1, value=int(o)).font = BOLD
        for j in range(2, 2 + len(grid.columns)):
            col = chr(64 + j)
            g.cell(row=i, column=j, value="=" + PAYBACK.format(o=f"$A{i}", v=f"{col}$4")).font = BLACK

    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None and cell.font.name != FONT:
                    f = cell.font
                    cell.font = Font(name=FONT, bold=f.bold, italic=f.italic, size=f.size, color=f.color)
    wb.calculation.fullCalcOnLoad = True
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    print("Saved", build())
